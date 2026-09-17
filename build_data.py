"""
Genera data.json para el tablero "Ausentismo y Deserción por Programa"
UNIMINUTO Sede Tolima-Huila, a partir del Excel institucional (SAP).

Formato de origen (una fila por CU + Programa, columnas en pares por periodo):
  Col A: Centro Universitario
  Col B: Descripción del programa
  Col C: Ausentes/Desertores 2014-1   Col D: % Ausentismo/Deserción 2014-1
  Col E: Ausentes/Desertores 2014-2   Col F: % Ausentismo/Deserción 2014-2
  ... y así sucesivamente un par de columnas por cada periodo.

Regla de vacíos (indicada por el usuario):
  - % en 0  -> sí hubo registro ese periodo, la tasa fue 0% (dato válido).
  - % vacío -> no hubo población ni ausentes/desertores registrados ese
               periodo para ese programa (sin dato, se guarda como null).
"""
import openpyxl, json, unicodedata, os

SRC = "/mnt/user-data/uploads/Historico_Ausentismo_y_Deserción_STH_-_Programas_31-8-2026.xlsx"

def _norm(s):
    return unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode().upper().strip()

# CU que no deben aparecer en el tablero de Programas (viven en el tablero de nivel CU)
EXCLUDED_CUS = {_norm(x) for x in ['Cajamarca', 'Florencia', 'Fresno', 'Líbano', 'Mariquita', 'Mocoa', 'Puerto Boyacá']}


def parse_sheet(ws):
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    # columnas 2,4,6... = conteo del periodo ; 3,5,7... = % del periodo
    periods = []
    for i in range(2, len(header), 2):
        label = (header[i] or '').split('\n')[0].strip()
        periods.append(label)

    programs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cu = row[0]
        if cu is None:
            continue
        cu = str(cu).strip()
        if _norm(cu) in EXCLUDED_CUS:
            continue
        programa = (row[1] or '').strip()
        if not programa:
            continue

        periods_dict = {}
        pct_series, idx_series = [], []
        for idx, per in enumerate(periods):
            cnt_cell = row[2 + 2 * idx]
            pct_cell = row[2 + 2 * idx + 1]
            if cnt_cell == '':
                cnt_cell = None
            if pct_cell == '':
                pct_cell = None
            if pct_cell is None:
                # sin registro ese periodo (sin población ni ausentes/desertores)
                periods_dict[per] = {"pct": None, "cnt": None, "pob": None}
                continue
            pct = round(pct_cell * 100, 2)
            cnt = int(cnt_cell) if cnt_cell is not None else 0
            pob = round(cnt / pct_cell) if pct_cell else None
            periods_dict[per] = {"pct": pct, "cnt": cnt, "pob": pob}
            pct_series.append(pct)
            idx_series.append(idx)

        active_sems = len(pct_series)
        avg_pct = round(sum(pct_series) / active_sems, 2) if active_sems else None

        slope = 0
        if active_sems >= 2:
            n = active_sems
            sx = sum(idx_series); sy = sum(pct_series)
            sxx = sum(x * x for x in idx_series); sxy = sum(x * y for x, y in zip(idx_series, pct_series))
            denom = (n * sxx - sx * sx)
            slope = round((n * sxy - sx * sy) / denom, 4) if denom else 0

        programs.append({
            "cu": cu, "programa": programa,
            "avg_pct": avg_pct, "trend_slope": slope, "active_sems": active_sems,
            "periods": periods_dict
        })

    return programs, periods


def build_cu_series(programs, periods):
    """Promedio simple (no ponderado por población, que no viene en el origen)
    de las tasas de los programas con dato válido en cada periodo, por CU y
    para el total de Sede (agregando todos los programas)."""
    cus = sorted({p["cu"] for p in programs})
    series = {}
    for cu in cus:
        cu_programs = [p for p in programs if p["cu"] == cu]
        vals_per_period = []
        for per in periods:
            vals = [p["periods"][per]["pct"] for p in cu_programs if p["periods"][per]["pct"] is not None]
            vals_per_period.append(round(sum(vals) / len(vals), 2) if vals else None)
        series[cu] = vals_per_period

    sede_vals = []
    for per in periods:
        vals = [p["periods"][per]["pct"] for p in programs if p["periods"][per]["pct"] is not None]
        sede_vals.append(round(sum(vals) / len(vals), 2) if vals else None)
    series["Sede Tolima-Huila"] = sede_vals
    return series


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    aus_programs, aus_periods = parse_sheet(wb['Ausentismo Programas'])
    des_programs, des_periods = parse_sheet(wb['Deserción Programas'])

    aus_cu_series = build_cu_series(aus_programs, aus_periods)
    des_cu_series = build_cu_series(des_programs, des_periods)

    data = {
        "ausentismo": {"periods": aus_periods, "cu_series": aus_cu_series, "programs": aus_programs},
        "desercion": {"periods": des_periods, "cu_series": des_cu_series, "programs": des_programs},
        "as_of": "31 de agosto de 2026"
    }

    with open("data.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    print("Ausentismo: programas", len(aus_programs), "periodos", len(aus_periods), aus_periods[0], "...", aus_periods[-1])
    print("Deserción: programas", len(des_programs), "periodos", len(des_periods), des_periods[0], "...", des_periods[-1])
    print("CUs:", sorted({p['cu'] for p in aus_programs} | {p['cu'] for p in des_programs}))
    print("data.json size:", round(os.path.getsize("data.json") / 1024, 1), "KB")


if __name__ == "__main__":
    main()
